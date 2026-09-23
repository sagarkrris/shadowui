from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

OUT = "/Users/sagkrish/shadowui/service-reconciler-registration-ejs-tests-troubleshooting.docx"

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
styles['Normal']._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
styles['Normal'].font.size = Pt(10.5)
styles['Normal'].paragraph_format.space_after = Pt(6)
for name, size in [('Title', 20), ('Heading 1', 14), ('Heading 2', 11.5)]:
    st = styles[name]
    st.font.name = 'Aptos Display' if name == 'Title' else 'Aptos'
    st._element.rPr.rFonts.set(qn('w:ascii'), st.font.name)
    st._element.rPr.rFonts.set(qn('w:hAnsi'), st.font.name)
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_before = Pt(12 if name != 'Title' else 0)
    st.paragraph_format.space_after = Pt(5)

# Remove the built-in Word title rule so the title is plain black typography.
title_ppr = styles['Title']._element.get_or_add_pPr()
for child in list(title_ppr):
    if child.tag == qn('w:pBdr'):
        title_ppr.remove(child)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def borders(cell, color='D9D9D9'):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in('w:tcBorders')
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders'); tcPr.append(tcBorders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = 'w:' + edge
        el = tcBorders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag); tcBorders.append(el)
        el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '6'); el.set(qn('w:color'), color)

def set_cell(cell, text, bold=False, white=False, size=9):
    cell.text = ''
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    r = p.add_run(text); r.bold = bold; r.font.size = Pt(size)
    if white: r.font.color.rgb = RGBColor(255,255,255)
    borders(cell)

def add_table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = 'Table Grid'
    for i, h in enumerate(headers):
        c=t.rows[0].cells[i]; shade(c, '17365D'); set_cell(c,h,True,True,9)
        c.width=Inches(widths[i])
    for n,row in enumerate(rows):
        cells=t.add_row().cells
        for i, value in enumerate(row):
            if n % 2 == 1: shade(cells[i], 'F2F6FA')
            set_cell(cells[i], value, size=8.6)
            cells[i].width=Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def bullet(text, level=0):
    p=doc.add_paragraph(style='List Bullet' if level==0 else 'List Bullet 2')
    p.add_run(text)
    return p

def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), rid)
    run = OxmlElement('w:r')
    rpr = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '0563C1'); rpr.append(color)
    underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single'); rpr.append(underline)
    run.append(rpr)
    text_el = OxmlElement('w:t'); text_el.text = text; run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)

title=doc.add_paragraph(style='Title')
title.alignment=WD_ALIGN_PARAGRAPH.LEFT
title.add_run('Service Reconciler Registration EJS Tests Troubleshooting')
p=doc.add_paragraph()
p.add_run('Purpose. ').bold=True
p.add_run('Use this runbook to interpret functional-test failures, verify environment prerequisites, and request help when deployment dependencies are missing.')
p=doc.add_paragraph()
p.add_run('Result summary. ').bold=True
p.add_run('The m18xf run recorded 6 errors and 2 failures, all in duplicate-alias queue verification. The m18bx run recorded 0 errors and 0 failures. The supplied results do not show an ArcticInterfaceNotFoundException for request 1215500; that exception has its own environment-resolution path below.')

doc.add_heading('Test run evidence', level=1)
add_table(['Result file', 'Target domain', 'Started tests', 'Errors', 'Failures', 'Outcome'], [
    ['java 20260922-174411.xml', 'm18xf.ip.devcerner.net', '140 of 142', '6', '2', 'Blocked alias queue checks'],
    ['java 20260922-202122.xml', 'm18bx.ip.devcerner.net', '140 of 142', '0', '0', 'Passed'],
], [1.55,1.42,0.78,0.55,0.62,1.58])

doc.add_heading('Failure and error inventory', level=1)
doc.add_paragraph('All eight affected scenarios are in the m18xf Registration Reconciler run. The details point to absent match-review or batch-combine queue data: lookup code calls Optional.get when no record is found, or the expected person identifier is null.')
add_table(['Type', 'Affected scenario', 'Evidence location', 'Observed result'], [
    ['Error (2)', 'Duplicate patient alias; add to batch combine queue', 'reconcile_aliases.feature:273 and :381; AliasesReconcilerSteps.java:449', 'NoSuchElementException: No value present'],
    ['Error (1)', 'Duplicate patient alias; add to match review queue', 'reconcile_aliases.feature:114; AliasesReconcilerSteps.java:419', 'NoSuchElementException: No value present'],
    ['Failure (1)', 'Duplicate patient alias; update existing match review queue record', 'reconcile_aliases.feature:216; AliasesReconcilerSteps.java:423', 'Expected 423456794, but was null'],
    ['Error (2)', 'Duplicate patient alias; add to batch combine queue', 'reconcile_patient_aliases.feature:273 and :381; PatientAliasesReconcilerSteps.java:448', 'NoSuchElementException: No value present'],
    ['Error (1)', 'Duplicate patient alias; add to match review queue', 'reconcile_patient_aliases.feature:114; PatientAliasesReconcilerSteps.java:418', 'NoSuchElementException: No value present'],
    ['Failure (1)', 'Duplicate patient alias; update existing match review queue record', 'reconcile_patient_aliases.feature:216; PatientAliasesReconcilerSteps.java:422', 'Expected 423456794, but was null'],
], [0.72,1.6,2.72,1.5])

doc.add_heading('How to handle the observed alias queue failures', level=1)
doc.add_paragraph('These failures are not direct evidence that request 1215500 is missing. First establish whether the reconciliation action created the expected queue record and whether the functional-test data is valid in m18xf.')
for text in [
    'Confirm the run target and test results. Compare against m18bx, where the same suite completed without errors.',
    'For the affected feature line, identify the matched people and the expected queue type: match review or batch combine.',
    'Query the relevant queue/test data in the target environment. Verify that the action created a record and that the returned person identifier is populated rather than null.',
    'Check test-data preconditions and cleanup. A missing, altered, historical, or previously consumed duplicate-alias fixture can prevent queue creation.',
    'If the record should exist but does not, collect the feature line, target domain, scenario name, request/response logs, and queue lookup result for the service owner. Do not change the assertion solely to hide missing environment data.',
    'Rerun only the affected scenarios after the environment or fixture is corrected. Then rerun the functional-test suite for that domain.'
]: bullet(text)

doc.add_heading('Alias configuration prerequisites', level=1)
doc.add_paragraph('The duplicate-alias scenarios require the following paired configuration. Confirm both systems use the same alias entity, alias type, and alias pool before rerunning. A mismatch can prevent the expected match-review or batch-combine queue record from being created.')
add_table(['System', 'Field or function', 'Match field', 'Alias entity', 'Alias type', 'Alias pool'], [
    ['Bedrock', 'Person Identifier', 'N/A', 'PERSON', 'Social Security Number', 'MSVC-ACCESS SSN (Dup N)'],
    ['Bedrock', 'Person Identifier', 'N/A', 'PERSON', 'Federal Person Principal', 'MSVC-ACCESS FEDPRINCIPAL (Auto, DEF)'],
    ['si_manager', 'Person Combine Match Queue', 'Alias', 'PERSON', 'Social Security Number', 'MSVC-ACCESS SSN (Dup N)'],
    ['si_manager', 'Person Combine Batch Queue', 'Alias', 'PERSON', 'Federal Person Principal', 'MSVC-ACCESS FEDPRINCIPAL (Auto, DEF)'],
], [0.72,1.22,0.72,0.72,1.28,1.75])
doc.add_paragraph('Verification order: confirm the two Bedrock alias pools first, then confirm the two si_manager match-function mappings. After correcting any missing or mismatched entry, rerun the affected duplicate-alias scenarios before running the full functional-test suite.')

doc.add_heading('Environment gate before functional tests', level=1)
doc.add_paragraph('Before running service-reconciler-registration-ejs-tests, verify that mlsystemintegrationconfig is deployed to the intended domain. The supplied Alva Watt dashboard screenshot shows deployments for m18xd, m18bx, and mbxl; it is the reference for this check.')
add_table(['Check', 'Pass condition', 'Action if not met'], [
    ['Alva Watt dashboard', 'Search for mlsystemintegrationconfig and find an entry for the target domain.', 'Do not rely on the domain for functional testing until the deployment is available.'],
    ['Deployment health', 'The matching deployment has a running pod and healthy status.', 'Wait for/resolve deployment health before test execution.'],
    ['TDB request 1215500', 'Request 1215500 is deployed to the target domain.', 'Use the collaboration-channel escalation below if the request is absent and the test call fails.'],
], [1.5,2.75,2.3])

doc.add_page_break()
doc.add_heading('Arctic interface error for request 1215500', level=1)
doc.add_paragraph('Use this path only when a functional test fails with ArcticInterfaceNotFoundException for request 1215500. It is an environment deployment dependency, distinct from the alias queue assertions recorded in the m18xf XML.')
for text in [
    'Capture the exact target domain and the full exception text.',
    'Verify in TDB whether request 1215500 is deployed to that domain.',
    'If it is deployed, verify mlsystemintegrationconfig deployment and pod health in the Alva Watt dashboard. Run the functional tests only after this prerequisite is present and healthy.',
    'If request 1215500 is not deployed, post the prepared message in #fsi-engineering-dev-environment-collab and wait for confirmation before rerunning.'
]: bullet(text)

doc.add_heading('Collaboration channel message', level=1)
p = doc.add_paragraph()
p.add_run('Reference thread. ').bold = True
add_hyperlink(p, 'Existing FSI Engineering Dev Environment Collaboration Slack thread', 'https://oracle-one.slack.com/archives/C06CD95RU68/p1774469882767059')
p=doc.add_paragraph()
p.paragraph_format.left_indent=Inches(0.2)
r=p.add_run('Hi, we are getting this error when calling request 1215500 in m18xb and m18xf:\n\ncom.cerner.system.enterprise.client.arctic.ArcticInterfaceNotFoundException: Call failure: Could not find the interface: name = 1215500, user = acsautotestuser, domain = m18xb.ip.devcerner.net\n\nand m18xf.ip.devcerner.net\n\nCan someone help with this issue?')
r.font.name='Aptos'; r.font.size=Pt(10)

doc.save(OUT)
print(OUT)
